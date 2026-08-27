"""Correct known GENE14 plate-map transcription gaps.

1. The 24 hpf crispant controls occupy P18 columns 1-2. Their hash-well formulas
   were missing for imaging columns 6-7.

2. The 48 hpf crispant embryos occupy two hash plates:

    imaging columns 1, 2, 3, 4, 5, 6, 7
    hash columns   10,11,12, 1, 2, 3, 4
    hash plate     18,18,18, 4, 4, 4, 4

3. The H row of cep290 30 hpf plate01 lost its hash-plate and RT-block values.
   Like rows A-G, it is an identity well map on P02 / Bl4.

Run once with the segmentation_grounded_sam Python environment. The script is
idempotent and verifies the saved workbook before exiting.
"""

from pathlib import Path

from openpyxl import load_workbook


HERE = Path(__file__).resolve().parent
CRISPANT_24_WORKBOOK = HERE / "20260319_cilia_crispant_24hpf_well_metadata.xlsx"
CRISPANT_48_WORKBOOK = HERE / "20260320_cilia_crispant_48hpf_well_metadata.xlsx"
CEP290_30_WORKBOOK = HERE / "20260324_cep290_30hpf_plate01_well_metadata.xlsx"

IMAGING_TO_HASH_COLUMN = {
    1: 10,
    2: 11,
    3: 12,
    4: 1,
    5: 2,
    6: 3,
    7: 4,
}

HASH_PLATE_BY_IMAGING_COLUMN = {
    1: 18,
    2: 18,
    3: 18,
    4: 4,
    5: 4,
    6: 4,
    7: 4,
}


def correct_crispant_24_map() -> None:
    workbook = load_workbook(CRISPANT_24_WORKBOOK)
    image_to_hash = workbook["image_to_hash_map"]

    for plate_row_number, row_letter in enumerate("ABCDEFGH", start=2):
        image_to_hash.cell(row=plate_row_number, column=7, value=f"{row_letter}01")
        image_to_hash.cell(row=plate_row_number, column=8, value=f"{row_letter}02")

    workbook.save(CRISPANT_24_WORKBOOK)


def verify_crispant_24_map() -> None:
    workbook = load_workbook(CRISPANT_24_WORKBOOK, data_only=True, read_only=True)
    image_to_hash = workbook["image_to_hash_map"]
    hash_plate = workbook["hash_plate_num"]
    rt_block = workbook["rt_block"]

    for plate_row_number, row_letter in enumerate("ABCDEFGH", start=2):
        for imaging_column, hash_column in ((6, 1), (7, 2)):
            excel_column = imaging_column + 1
            if image_to_hash.cell(plate_row_number, excel_column).value != f"{row_letter}0{hash_column}":
                raise ValueError(f"{row_letter}{imaging_column}: wrong 24 hpf hash well")
            if hash_plate.cell(plate_row_number, excel_column).value != 18:
                raise ValueError(f"{row_letter}{imaging_column}: expected hash plate 18")
            if rt_block.cell(plate_row_number, excel_column).value != "Bl1":
                raise ValueError(f"{row_letter}{imaging_column}: expected RT block Bl1")


def correct_crispant_48_map() -> None:
    workbook = load_workbook(CRISPANT_48_WORKBOOK)
    image_to_hash = workbook["image_to_hash_map"]
    hash_plate = workbook["hash_plate_num"]

    # Excel row 1 and column A are headers. Plate rows A-H occupy Excel rows 2-9,
    # and imaging columns 1-7 occupy Excel columns B-H.
    for plate_row_number, row_letter in enumerate("ABCDEFGH", start=2):
        for imaging_column in range(1, 8):
            excel_column = imaging_column + 1
            hash_column = IMAGING_TO_HASH_COLUMN[imaging_column]

            image_to_hash.cell(
                row=plate_row_number,
                column=excel_column,
                value=f"{row_letter}{hash_column:02d}",
            )
            hash_plate.cell(
                row=plate_row_number,
                column=excel_column,
                value=HASH_PLATE_BY_IMAGING_COLUMN[imaging_column],
            )

    workbook.save(CRISPANT_48_WORKBOOK)


def verify_crispant_48_map() -> None:
    workbook = load_workbook(CRISPANT_48_WORKBOOK, data_only=True, read_only=True)
    image_to_hash = workbook["image_to_hash_map"]
    hash_plate = workbook["hash_plate_num"]

    for plate_row_number, row_letter in enumerate("ABCDEFGH", start=2):
        for imaging_column in range(1, 8):
            excel_column = imaging_column + 1
            expected_hash_column = IMAGING_TO_HASH_COLUMN[imaging_column]
            expected_hash_well = f"{row_letter}{expected_hash_column:02d}"
            expected_hash_plate = HASH_PLATE_BY_IMAGING_COLUMN[imaging_column]

            actual_hash_well = image_to_hash.cell(plate_row_number, excel_column).value
            actual_hash_plate = hash_plate.cell(plate_row_number, excel_column).value

            if actual_hash_well != expected_hash_well:
                raise ValueError(
                    f"{row_letter}{imaging_column}: expected hash well "
                    f"{expected_hash_well}, found {actual_hash_well}"
                )
            if actual_hash_plate != expected_hash_plate:
                raise ValueError(
                    f"{row_letter}{imaging_column}: expected hash plate "
                    f"{expected_hash_plate}, found {actual_hash_plate}"
                )


def correct_cep290_30_h_row() -> None:
    workbook = load_workbook(CEP290_30_WORKBOOK)
    hash_plate = workbook["hash_plate_num"]
    rt_block = workbook["rt_block"]

    # Plate row H is Excel row 9. All twelve wells are identity-mapped to P02 / Bl4.
    for imaging_column in range(1, 13):
        excel_column = imaging_column + 1
        hash_plate.cell(row=9, column=excel_column, value=2)
        rt_block.cell(row=9, column=excel_column, value="Bl4")

    workbook.save(CEP290_30_WORKBOOK)


def verify_cep290_30_h_row() -> None:
    workbook = load_workbook(CEP290_30_WORKBOOK, data_only=True, read_only=True)
    image_to_hash = workbook["image_to_hash_map"]
    hash_plate = workbook["hash_plate_num"]
    rt_block = workbook["rt_block"]

    for imaging_column in range(1, 13):
        excel_column = imaging_column + 1
        expected_hash_well = f"H{imaging_column:02d}"

        if image_to_hash.cell(9, excel_column).value != expected_hash_well:
            raise ValueError(f"H{imaging_column}: expected identity hash well {expected_hash_well}")
        if hash_plate.cell(9, excel_column).value != 2:
            raise ValueError(f"H{imaging_column}: expected hash plate 2")
        if rt_block.cell(9, excel_column).value != "Bl4":
            raise ValueError(f"H{imaging_column}: expected RT block Bl4")


if __name__ == "__main__":
    correct_crispant_24_map()
    verify_crispant_24_map()
    print(f"corrected and verified {CRISPANT_24_WORKBOOK}")

    correct_crispant_48_map()
    verify_crispant_48_map()
    print(f"corrected and verified {CRISPANT_48_WORKBOOK}")

    correct_cep290_30_h_row()
    verify_cep290_30_h_row()
    print(f"corrected and verified {CEP290_30_WORKBOOK}")
