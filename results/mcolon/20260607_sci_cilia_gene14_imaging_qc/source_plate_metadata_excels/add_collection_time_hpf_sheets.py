"""Add a collection_time_hpf plate to the 24 GENE14 metadata workbooks.

Most experiments have one collection time. The mixed 18/24 hpf experiment is
still simple: Bl2 was collected at 18 hpf and Bl3 was collected at 24 hpf.
"""

from pathlib import Path

from openpyxl import load_workbook


HERE = Path(__file__).resolve().parent

COLLECTION_TIME_BY_FILE = {
    "20260319_cilia_crispant_18hpf_well_metadata.xlsx": 18,
    "20260319_cilia_crispant_24hpf_well_metadata.xlsx": 24,
    "20260319_cilia_crispant_30hpf_well_metadata.xlsx": 30,
    "20260320_cilia_crispant_48hpf_well_metadata.xlsx": 48,
    "20260324_cep290_18hpf_plate01_well_metadata.xlsx": 18,
    "20260324_cep290_24hpf_plate01_well_metadata.xlsx": 24,
    "20260324_cep290_24hpf_plate02_well_metadata.xlsx": 24,
    "20260324_cep290_30hpf_plate01_well_metadata.xlsx": 30,
    "20260324_cep290_30hpf_plate02_well_metadata.xlsx": 30,
    "20260331_b9d2_18hpf_plate01_well_metadata.xlsx": 18,
    "20260331_b9d2_18hpf_plate02_well_metadata.xlsx": 18,
    "20260414_b9d2_14hpf_plate01_well_metadata.xlsx": 14,
    "20260414_b9d2_14hpf_plate02_well_metadata.xlsx": 14,
    "20260414_b9d2_30hpf_plate01_well_metadata.xlsx": 30,
    "20260414_b9d2_30hpf_plate02_well_metadata.xlsx": 30,
    "20260414_sci_b9d2_48hpf_plate01_well_metadata.xlsx": 48,
    "20260415_b9d2_30to48hpf_plate01_t02_well_metadata.xlsx": 48,
    "20260415_b9d2_30to48hpf_plate02_t02_well_metadata.xlsx": 48,
    "20260415_cep290_18hpf_plate03_well_metadata.xlsx": 18,
    "20260415_cep290_30to48hpf_plate02_t01_well_metadata.xlsx": 30,
    "20260415_sci_cep290_48hpf_plate01_well_metadata.xlsx": 48,
    "20260416_cep290_30to48hpf_plate01_t02_well_metadata.xlsx": 48,
    "20260416_cep290_30to48hpf_plate02_t02_well_metadata.xlsx": 48,
}

MIXED_FILE = "20260324_cep290_18hpf_24hpf_plate02_well_metadata.xlsx"
COLLECTION_TIME_BY_RT_BLOCK = {"Bl2": 18, "Bl3": 24}


def add_collection_time_sheet(path: Path, constant_time: int | None = None) -> None:
    workbook = load_workbook(path)

    # Copying the RT-block plate gives the new sheet the same tidy 8 x 12 layout.
    if "collection_time_hpf" in workbook.sheetnames:
        del workbook["collection_time_hpf"]
    collection_time = workbook.copy_worksheet(workbook["rt_block"])
    collection_time.title = "collection_time_hpf"
    rt_block = workbook["rt_block"]

    for row in range(2, 10):
        for column in range(2, 14):
            block = rt_block.cell(row, column).value
            if block is None:
                value = None
            elif constant_time is not None:
                value = constant_time
            else:
                value = COLLECTION_TIME_BY_RT_BLOCK.get(block)
                if value is None:
                    raise ValueError(f"{path.name}: no collection time for RT block {block!r}")

            collection_time.cell(row, column, value)

    workbook.save(path)


def verify_collection_time_sheet(path: Path, constant_time: int | None = None) -> None:
    workbook = load_workbook(path, data_only=True, read_only=True)
    rt_block = workbook["rt_block"]
    collection_time = workbook["collection_time_hpf"]

    for row in range(2, 10):
        for column in range(2, 14):
            block = rt_block.cell(row, column).value
            expected = (
                None
                if block is None
                else constant_time
                if constant_time is not None
                else COLLECTION_TIME_BY_RT_BLOCK[block]
            )
            actual = collection_time.cell(row, column).value
            if actual != expected:
                raise ValueError(
                    f"{path.name} cell {collection_time.cell(row, column).coordinate}: "
                    f"expected {expected}, found {actual}"
                )


if __name__ == "__main__":
    for filename, collection_time in COLLECTION_TIME_BY_FILE.items():
        workbook_path = HERE / filename
        add_collection_time_sheet(workbook_path, collection_time)
        verify_collection_time_sheet(workbook_path, collection_time)
        print(f"{filename}: {collection_time} hpf")

    mixed_path = HERE / MIXED_FILE
    add_collection_time_sheet(mixed_path)
    verify_collection_time_sheet(mixed_path)
    print(f"{MIXED_FILE}: Bl2 -> 18 hpf, Bl3 -> 24 hpf")
