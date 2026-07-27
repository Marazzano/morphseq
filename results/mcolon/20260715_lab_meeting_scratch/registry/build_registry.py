#!/usr/bin/env python
"""Build a registry of experiments containing cilia-mutant genotypes.

Genes tracked (substring match on the gene token, case-insensitive):
    cep290, tmem67, rpgrip1l, b9d2

Sources (priority order):
  1. Plate metadata genotype sheets: metadata/plate_metadata/*.xlsx, sheet
     named 'genotype' (case-insensitive). Files with '_backup_' or '_fixed'
     in the name are skipped. Every cell of the genotype sheet is scanned;
     any cell containing a gene token marks that experiment for that gene.
     experiment_id = leading date token of the filename (first '_'-split field).
  2. build06 output: morphseq_playground/metadata/build06_output/
     df03_final_output_with_latents_*.csv. Provides embryo-level detail:
     unique embryo_id (use_embryo_flag=True) and predicted_stage_hpf min/max,
     per gene, matched on the genotype column.

Output: cilia_mutant_experiment_registry.csv (one row per experiment_id x gene).

Refresh: rerun this script (rebuilds the CSV from current files).

Read-only w.r.t. all sources; writes only into this registry/ directory.
"""
from __future__ import annotations
import glob
import os
import re
import sys
import pandas as pd
import openpyxl

REPO = "/net/trapnell/vol1/home/mdcolon/proj/morphseq"
PLATE_DIR = os.path.join(REPO, "metadata", "plate_metadata")
BUILD06_DIR = os.path.join(REPO, "morphseq_playground", "metadata", "build06_output")
OUT_DIR = os.path.dirname(os.path.abspath(__file__))
OUT_CSV = os.path.join(OUT_DIR, "cilia_mutant_experiment_registry.csv")

GENES = ["cep290", "tmem67", "rpgrip1l", "b9d2"]
# Match the exact gene token. rpgrip1l must not be triggered by 'rpgrip1'
# (rpgrip1l already contains rpgrip1, so a plain substring for rpgrip1l is
# safe; we simply never search for a bare 'rpgrip1' token).


def gene_in_text(text: str, gene: str) -> bool:
    return gene in text.lower()


def scan_plate_metadata():
    """Return {experiment_id: {gene: filename}} and the token forms seen."""
    result = {}  # exp_id -> {gene: basename}
    tokens_seen = {g: set() for g in GENES}
    files = sorted(glob.glob(os.path.join(PLATE_DIR, "*.xlsx")))
    n_scanned = 0
    for path in files:
        base = os.path.basename(path)
        if "_backup_" in base or "_fixed" in base:
            continue
        exp_id = base.split("_")[0]
        n_scanned += 1
        try:
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        except Exception as e:
            print(f"  [warn] cannot open {base}: {e}", file=sys.stderr)
            continue
        sheet_name = next((s for s in wb.sheetnames if s.lower() == "genotype"), None)
        if sheet_name is None:
            wb.close()
            continue
        ws = wb[sheet_name]
        cell_texts = []
        for row in ws.iter_rows(values_only=True):
            for v in row:
                if v is not None:
                    cell_texts.append(str(v))
        wb.close()
        blob = " ".join(cell_texts).lower()
        for g in GENES:
            if g in blob:
                result.setdefault(exp_id, {})[g] = base
                # record token forms (word-ish tokens containing the gene)
                for tok in re.findall(r"[a-z0-9_]+", blob):
                    if g in tok:
                        tokens_seen[g].add(tok)
    print(f"  scanned {n_scanned} plate-metadata workbooks", file=sys.stderr)
    return result, tokens_seen


def scan_build06():
    """Return {experiment_id: {gene: {n_embryos, stage_min, stage_max}}}."""
    result = {}
    files = sorted(glob.glob(os.path.join(BUILD06_DIR, "df03_final_output_with_latents_*.csv")))
    for path in files:
        base = os.path.basename(path)
        if base.endswith(".archive.csv"):
            continue
        try:
            df = pd.read_csv(path, usecols=lambda c: c in (
                "experiment_id", "embryo_id", "genotype",
                "predicted_stage_hpf", "use_embryo_flag"), low_memory=False)
        except Exception as e:
            print(f"  [warn] cannot read {base}: {e}", file=sys.stderr)
            continue
        if "genotype" not in df.columns or "experiment_id" not in df.columns:
            continue
        df = df[df.get("use_embryo_flag", True) == True] if "use_embryo_flag" in df.columns else df
        if df.empty:
            continue
        df = df.copy()
        df["genotype"] = df["genotype"].astype(str).str.lower()
        for g in GENES:
            sub = df[df["genotype"].str.contains(g, na=False)]
            if sub.empty:
                continue
            for exp_id, grp in sub.groupby("experiment_id"):
                exp_id = str(exp_id)
                n_emb = grp["embryo_id"].nunique() if "embryo_id" in grp else len(grp)
                smin = smax = None
                if "predicted_stage_hpf" in grp:
                    st = pd.to_numeric(grp["predicted_stage_hpf"], errors="coerce").dropna()
                    if len(st):
                        smin, smax = float(st.min()), float(st.max())
                rec = result.setdefault(exp_id, {}).setdefault(
                    g, {"n_embryos": 0, "stage_min": None, "stage_max": None})
                rec["n_embryos"] += n_emb
                if smin is not None:
                    rec["stage_min"] = smin if rec["stage_min"] is None else min(rec["stage_min"], smin)
                    rec["stage_max"] = smax if rec["stage_max"] is None else max(rec["stage_max"], smax)
    return result


def main():
    print("Scanning plate metadata ...", file=sys.stderr)
    plate, tokens = scan_plate_metadata()
    print("Scanning build06 output ...", file=sys.stderr)
    b06 = scan_build06()

    # union of all (exp, gene) keys
    keys = set()
    for exp, genes in plate.items():
        for g in genes:
            keys.add((exp, g))
    for exp, genes in b06.items():
        for g in genes:
            keys.add((exp, g))

    rows = []
    for exp, g in sorted(keys):
        in_plate = exp in plate and g in plate[exp]
        in_b06 = exp in b06 and g in b06[exp]
        plate_file = plate.get(exp, {}).get(g, "")
        n_emb = stage_min = stage_max = ""
        notes = []
        if in_b06:
            r = b06[exp][g]
            n_emb = r["n_embryos"]
            stage_min = "" if r["stage_min"] is None else round(r["stage_min"], 2)
            stage_max = "" if r["stage_max"] is None else round(r["stage_max"], 2)
        if in_plate and not in_b06:
            notes.append("metadata-only (no build06)")
        if in_b06 and not in_plate:
            notes.append("build06-only (no plate-metadata genotype sheet)")
        rows.append({
            "experiment_id": exp,
            "gene": g,
            "in_plate_metadata": in_plate,
            "in_build06": in_b06,
            "n_embryos": n_emb,
            "stage_min_hpf": stage_min,
            "stage_max_hpf": stage_max,
            "plate_metadata_file": plate_file,
            "notes": "; ".join(notes),
        })

    out = pd.DataFrame(rows, columns=[
        "experiment_id", "gene", "in_plate_metadata", "in_build06",
        "n_embryos", "stage_min_hpf", "stage_max_hpf",
        "plate_metadata_file", "notes"])
    out = out.sort_values(["experiment_id", "gene"]).reset_index(drop=True)
    out.to_csv(OUT_CSV, index=False)

    # summary
    print("\n=== Token forms matched (plate metadata) ===", file=sys.stderr)
    for g in GENES:
        print(f"  {g}: {sorted(tokens[g])}", file=sys.stderr)
    print("\n=== Experiments per gene ===", file=sys.stderr)
    for g in GENES:
        n = out[out["gene"] == g]["experiment_id"].nunique()
        print(f"  {g}: {n}", file=sys.stderr)
    print(f"\nTotal unique experiments: {out['experiment_id'].nunique()}", file=sys.stderr)
    print(f"Total rows: {len(out)}", file=sys.stderr)
    print(f"Wrote {OUT_CSV}", file=sys.stderr)
    return out


if __name__ == "__main__":
    main()
