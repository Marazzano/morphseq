#!/usr/bin/env python
"""Refresh the Keyence/YX1 pipeline audit workbook and executable dashboard notebook."""

from __future__ import annotations

import argparse
import json
import math
import re
import statistics
import subprocess
import sys
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Iterable

import nbformat as nbf
import pandas as pd
from openpyxl import load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


AUDIT_DIR = Path(__file__).resolve().parent
REPO_ROOT = Path(__file__).resolve().parents[3]
PIPELINE_ROOT = Path(
    "/net/trapnell/vol1/home/nlammers/projects/data/morphseq/pipeline/output"
)
MANIFEST = (
    REPO_ROOT
    / "src/data_pipeline/pipeline_orchestrator/manifests/front_half_archive.txt"
)
LOG_DIR = REPO_ROOT / "logs"
WORKBOOK = AUDIT_DIR / "pipeline_audit.xlsx"
NOTEBOOK = AUDIT_DIR / "pipeline_failure_dashboard.ipynb"


@dataclass(frozen=True)
class Failure:
    category: str
    subcategory: str
    stage: str
    effort_tier: int
    remediation_state: str
    error_signature: str
    solution: str


STAGES = [
    ("metadata_ingest", "Metadata ingest"),
    ("materialization", "Materialization / frame inventory"),
    ("frame_detections", "GroundingDINO detection"),
    ("frame_masks", "SAM2 segmentation"),
    ("physical_embryo_registry", "Physical embryo registry"),
    ("snip_processing", "Snip processing"),
    ("focus_qc", "Focus QC"),
    ("motion_blur_qc", "Motion-blur QC"),
    ("mask_quality_qc", "Mask-quality QC"),
    ("surface_area_qc", "Surface-area QC"),
    ("snip_qc", "Combined snip QC"),
    ("death_detection", "Death detection"),
    ("mask_geometry", "Mask geometry"),
    ("curvature_metrics", "Curvature metrics"),
    ("stage_predictions", "Stage predictions"),
    ("latent_embeddings", "Latent embeddings"),
    ("pose_kinematics", "Pose/kinematics"),
    ("fraction_alive", "Fraction alive"),
    ("analysis_ready", "Analysis-ready"),
]
STAGE_ORDER = {name: i for i, (name, _) in enumerate(STAGES)}


def _failure(
    category: str,
    subcategory: str,
    stage: str,
    effort_tier: int,
    remediation_state: str,
    error_signature: str,
    solution: str,
) -> Failure:
    return Failure(
        category,
        subcategory,
        stage,
        effort_tier,
        remediation_state,
        error_signature,
        solution,
    )


def known_failures() -> dict[str, list[Failure]]:
    out: dict[str, list[Failure]] = {}

    def add(experiments: Iterable[str], issue: Failure) -> None:
        for experiment in experiments:
            out.setdefault(experiment, []).append(issue)

    add(
        ["20230830", "20231207", "20231208", "20240507"],
        _failure(
            "Input eligibility",
            "Abandoned/test paths included by recursive discovery",
            "metadata_ingest",
            2,
            "code_change_needed",
            "Unexpected channel or duplicate inventory rows from ignore/other_files",
            "Exclude path components such as ignore and other_files in one centralized "
            "Keyence source-eligibility policy.",
        ),
    )
    add(
        ["20230608"],
        _failure(
            "Source corruption",
            "Single zero-byte TIFF plane",
            "metadata_ingest",
            3,
            "data_recovery_needed",
            "No <Data> block in TIFF suffix; inspected TIFF is zero bytes",
            "Quarantine the unavailable z-plane, retain remaining planes, and allow projection "
            "when the residual stack is sufficient.",
        ),
    )
    add(
        [
            "20250529_24hpf_ctrl_atf6",
            "20250529_24hpf_wfs1_ctcf",
            "20250529_30hpf_ctrl_atf6",
            "20250529_30hpf_wfs1_ctcf",
            "20250529_36hpf_ctrl_atf6",
            "20250529_36hpf_wfs1_ctcf",
        ],
        _failure(
            "Materialization geometry",
            "Stitch master shape/size mismatch",
            "materialization",
            2,
            "code_change_needed",
            "Master stitch parameters are 1x6; current mosaic is 1x3 (size remains 6)",
            "Normalize both metadata.shape and metadata.size before reloading master stitch "
            "parameters.",
        ),
    )
    add(
        ["20240314"],
        _failure(
            "Metadata dialect",
            "Unmapped YX1 fluorescence channel",
            "metadata_ingest",
            2,
            "code_change_needed",
            "No canonical mapping for raw channel 'Celesta 473'",
            "Add the explicit tested alias Celesta 473 -> GFP.",
        ),
    )
    add(
        ["20250520"],
        _failure(
            "Truncated acquisition",
            "ND2 padded with a zero-valued time-series tail",
            "materialization",
            3,
            "data_recovery_needed",
            "Focus-stack intensity bounds collapse to lo=hi=0 after acquisition stops",
            "Trim each position's contiguous zero-valued terminal suffix and retain the 3,728 "
            "valid well-timepoints without padding.",
        ),
    )
    add(
        ["20260124"],
        _failure(
            "Orchestration race",
            "Shared merged runtime config is written non-atomically",
            "metadata_ingest",
            2,
            "code_change_needed",
            "PyYAML ScannerError while reading truncated merged_config.yaml",
            "Use a per-experiment/run config path and finalize writes atomically with os.replace.",
        ),
    )
    add(
        [
            "20240813_extras",
            "20260702_hotchem_24hpf_plate01",
            "20260702_hotchem_24hpf_plate02",
            "20260702_hotchem_30hpf_plate01",
            "20260702_hotchem_30hpf_plate02",
        ],
        _failure(
            "Workflow state",
            "Stale Snakemake lock",
            "frame_detections",
            1,
            "operational_cleanup",
            "Directory cannot be locked / stale Snakemake lock",
            "Confirm no live owner, run Snakemake --unlock for the experiment work directory, "
            "then requeue.",
        ),
    )
    add(
        [
            "20250612_24hpf_ctrl_atf6",
            "20250612_24hpf_wfs1_ctcf",
            "20250612_30hpf_ctrl_atf6",
            "20250612_30hpf_wfs1_ctcf",
            "20260320_cilia_crispant_48hpf",
            "20260324_cep290_18hpf_24hpf_plate02",
            "20260324_cep290_18hpf_plate01",
            "20260324_cep290_24hpf_plate01",
            "20260324_cep290_24hpf_plate02",
            "20260324_cep290_30hpf_plate01",
            "20260331_b9d2_18hpf_plate02",
            "20260414_b9d2_14hpf_plate01",
            "20260414_b9d2_14hpf_plate02",
            "20260414_b9d2_30hpf_plate02",
            "20260415_b9d2_30to48hpf_plate02_t02",
            "20260415_cep290_18hpf_plate03",
            "20260415_cep290_30to48hpf_plate02_t01",
        ],
        _failure(
            "Materialization geometry",
            "Per-z-plane FLANN alignment",
            "materialization",
            1,
            "fixed_in_head_needs_rerun",
            "OpenCV FLANN: (size_t)knn <= index_->size()",
            "Current code aligns once on the focus-stacked frame and reuses transforms. Force "
            "regeneration of affected Keyence z-stack products and rerun.",
        ),
    )
    add(
        [
            "20260324_cep290_18hpf_24hpf_plate02",
            "20260324_cep290_18hpf_plate01",
            "20260331_b9d2_18hpf_plate01",
            "20260414_b9d2_14hpf_plate02",
        ],
        _failure(
            "Acquisition reconciliation",
            "Same well was reacquired under a second XY position",
            "materialization",
            3,
            "data_reconciliation_needed",
            "Duplicate time/tile/z inventory cells",
            "Choose the most complete acquisition, then the latest capture as tie-breaker; keep "
            "the losing source only in an audit artifact.",
        ),
    )
    add(
        [
            "20260320_cilia_crispant_48hpf",
            "20260324_cep290_24hpf_plate02",
            "20260414_b9d2_30hpf_plate02",
            "20260415_b9d2_30to48hpf_plate02_t02",
            "20260415_cep290_18hpf_plate03",
            "20260415_cep290_30to48hpf_plate02_t01",
        ],
        _failure(
            "Developmental metadata",
            "Missing start_age_hpf",
            "stage_predictions",
            1,
            "fixed_in_head_needs_rerun",
            "start_age_hpf is required for stage prediction",
            "Current code emits null predicted stage with missing_start_age_hpf status. Force "
            "stage/QC regeneration.",
        ),
    )
    add(
        ["20250126"],
        _failure(
            "Developmental metadata",
            "Missing temperature",
            "stage_predictions",
            1,
            "fixed_in_head_needs_rerun",
            "temperature is required for stage prediction",
            "Current code emits null predicted stage with missing_temperature status. Force "
            "stage/QC regeneration.",
        ),
    )
    add(
        ["20260320_cilia_crispant_48hpf"],
        _failure(
            "Detection/segmentation",
            "No retained prompt detections in well A10",
            "frame_detections",
            3,
            "segmentation_retry_needed",
            "prompt_detections must contain at least one kept row",
            "Retry detection with a controlled fallback; if it still fails, quarantine only A10 "
            "with an explicit no_valid_detection status.",
        ),
    )
    add(
        ["20250215", "20250305", "20250415", "20250416", "20250425"],
        _failure(
            "Artifact compatibility",
            "Pre-refactor QC/stage tables have stale schemas",
            "focus_qc",
            1,
            "rerun_needed",
            "Required applicability/status columns are absent from existing tables",
            "Invalidate and rebuild affected stage/QC products; add schema-version dependency for "
            "future automatic invalidation.",
        ),
    )
    add(
        ["20251017_part2", "20251020"],
        _failure(
            "Metadata dialect",
            "Blank eighth plate-row label",
            "metadata_ingest",
            2,
            "code_change_needed",
            "Observed plate rows A-G, NaN; expected A-H",
            "When and only when there are exactly eight rows with A-G followed by one blank, "
            "canonicalize the last label to H.",
        ),
    )
    add(
        [
            "20260320",
            "20250416",
            "20251104",
            "20251106",
            "20251113",
            "20260223",
            "20260224",
        ],
        _failure(
            "Detection/segmentation",
            "All masks invalid; snip writer emits a headerless empty CSV",
            "snip_processing",
            3,
            "segmentation_retry_needed",
            "pandas.errors.EmptyDataError: No columns to parse from file",
            "Always emit a schemaful empty table and quarantine the failed well; separately retry "
            "detection/segmentation to recover its biological observations.",
        ),
    )
    add(
        [
            "20260320",
            "20251106",
            "20251121",
            "20251125",
            "20260213",
            "20260219",
            "20260319",
        ],
        _failure(
            "Join semantics",
            "Death-stage lookup is unnecessarily embryo-specific",
            "death_detection",
            2,
            "code_change_needed",
            "death_event: no stage_predictions row for physical_embryo_id at called-death frame",
            "Calculate stage from well/time acquisition metadata independently of embryo/snip "
            "presence; preserve null/status behavior for missing age or temperature.",
        ),
    )
    add(
        ["20260304"],
        _failure(
            "QC robustness",
            "Motion-blur QC receives a zero-pixel aligned mask",
            "motion_blur_qc",
            3,
            "well_quarantine_needed",
            "motion_blur_qc: aligned mask contains zero pixels",
            "Mark motion-blur QC not applicable for that snip/well, quarantine the invalid mask, "
            "and retry upstream segmentation if point recovery is required.",
        ),
    )
    add(
        [
            "20260326_wt_ref",
            "20260410_otx_pilot",
            "20260414_sci_b9d2_48hpf_plate01",
            "20260415_sci_cep290_48hpf_plate01",
            "20260501_zfpm_pilot",
            "20260502_zfpm_pilot",
            "20260417_irx_pilot",
            "20260418_irx_pilot",
        ],
        _failure(
            "Filesystem permissions",
            "Snakemake work directory is not writable",
            "frame_detections",
            1,
            "operational_cleanup",
            "PermissionError below work_directories/<experiment>/.snakemake",
            "Use a work-directory root owned by nlammers or repair group ownership/write "
            "permissions, then requeue.",
        ),
    )
    return out


def parse_manifest() -> pd.DataFrame:
    rows = []
    for line in MANIFEST.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line or line.startswith("#"):
            continue
        experiment_id, scope = line.split()[:2]
        rows.append({"experiment_id": experiment_id, "scope": scope})
    return pd.DataFrame(rows)


def read_tail(path: Path, max_bytes: int = 524_288) -> str:
    if not path.is_file():
        return ""
    with path.open("rb") as handle:
        size = path.stat().st_size
        handle.seek(max(0, size - max_bytes))
        return handle.read().decode("utf-8", errors="replace")


def extract_experiment(text: str) -> str | None:
    patterns = [
        r"(?m)^\s*EXPERIMENT\s*:\s*(\S+)",
        r"(?m)^experiment=(\S+)",
        r"(?m)^Experiment:\s*(\S+)",
        r"Experiments:\s*\['([^']+)'\]",
    ]
    for pattern in patterns:
        match = re.search(pattern, text)
        if match:
            return match.group(1).strip()
    return None


def extract_error_signature(error_tail: str) -> str:
    candidates = []
    patterns = (
        r"(?m)^(?:[A-Za-z_][\w.]*\.)?(?:ValueError|RuntimeError|PermissionError|"
        r"FileNotFoundError|AssertionError|EmptyDataError|ScannerError): .+$",
        r"(?m)^ERROR: .+$",
        r"(?m)^Error: Directory cannot be locked.+$",
    )
    for pattern in patterns:
        candidates.extend(re.findall(pattern, error_tail))
    if not candidates:
        return ""
    return re.sub(r"\s+", " ", candidates[-1]).strip()[:500]


def collect_logs() -> pd.DataFrame:
    rows = []
    pattern = re.compile(r"^(front_half|back_half)\.(\d+)\.(\d+)\.out$")
    for out_path in sorted(LOG_DIR.glob("*_half.*.*.out")):
        match = pattern.match(out_path.name)
        if not match:
            continue
        run_kind, job_id, task_id = match.groups()
        text = out_path.read_text(encoding="utf-8", errors="replace")
        experiment_id = extract_experiment(text)
        if not experiment_id:
            continue
        err_path = out_path.with_suffix(".err")
        err_tail = read_tail(err_path)
        finished = bool(re.search(r"(?:front_half|back_half) FINISHED", text))
        failed = bool(
            re.search(
                r"Shutting down, this might take some time|"
                r"Exiting because a job execution failed|"
                r"PermissionError:|ERROR: stale Snakemake lock|"
                r"Error: Directory cannot be locked",
                err_tail + "\n" + text[-20_000:],
            )
        )
        mtime = max(
            out_path.stat().st_mtime,
            err_path.stat().st_mtime if err_path.exists() else 0,
        )
        if finished:
            state = "success"
        elif failed:
            state = "failed"
        elif datetime.now().timestamp() - mtime < 8 * 3600:
            state = "running_or_incomplete"
        else:
            state = "incomplete_unknown"
        rows.append(
            {
                "experiment_id": experiment_id,
                "run_kind": run_kind,
                "job_id": int(job_id),
                "task_id": int(task_id),
                "run_state": state,
                "inferred_exit_status": 0 if finished else (1 if failed else pd.NA),
                "error_signature": extract_error_signature(err_tail),
                "stdout_path": str(out_path),
                "stderr_path": str(err_path) if err_path.exists() else "",
                "log_mtime": datetime.fromtimestamp(mtime),
            }
        )
    if not rows:
        return pd.DataFrame()
    return pd.DataFrame(rows).sort_values(["experiment_id", "log_mtime"])


def latest_log(logs: pd.DataFrame, experiment_id: str, run_kind: str) -> dict:
    if logs.empty:
        return {}
    selected = logs[
        (logs["experiment_id"] == experiment_id) & (logs["run_kind"] == run_kind)
    ]
    if selected.empty:
        return {}
    return selected.iloc[-1].to_dict()


def count_lines(path: Path) -> int:
    with path.open("rb") as handle:
        return max(0, sum(1 for _ in handle) - 1)


def approximate_frames(frame_files: list[Path]) -> tuple[int | None, str]:
    if not frame_files:
        return None, "No canonical frame inventories available"
    sample = frame_files[: min(5, len(frame_files))]
    counts = [count_lines(path) for path in sample]
    estimate = int(round(statistics.median(counts) * len(frame_files)))
    return estimate, (
        f"Median rows from {len(sample)} sampled per-well canonical frame inventories "
        f"x {len(frame_files)} wells"
    )


def project_name(experiment_id: str) -> str:
    low = experiment_id.lower()
    rules = [
        ("hotchem", "Chemical perturbation"),
        ("chem", "Chemical perturbation"),
        ("cep290", "CEP290"),
        ("b9d2", "B9D2"),
        ("cilia", "Cilia crispant"),
        ("irx", "IRX"),
        ("zfpm", "ZFPM"),
        ("otx", "OTX"),
        ("pbx", "PBX"),
        ("wt_ref", "Wild-type reference"),
        ("atf6", "ER-stress regulators"),
        ("wfs1", "ER-stress regulators"),
        ("ctcf", "ER-stress regulators"),
    ]
    for token, label in rules:
        if token in low:
            return label
    return "Legacy / unspecified"


def metadata_hints(experiment_id: str) -> tuple[str, str, int | None]:
    plate = (
        PIPELINE_ROOT
        / "acquisition"
        / experiment_id
        / "ingest_metadata"
        / "plate_metadata.csv"
    )
    if not plate.is_file():
        return "", "", None
    try:
        frame = pd.read_csv(plate, nrows=192)
    except Exception:
        return "", "", None

    def summarize(column: str) -> str:
        if column not in frame:
            return ""
        values = (
            frame[column]
            .dropna()
            .astype(str)
            .str.strip()
            .loc[lambda series: ~series.isin(["", "nan", "none", "not recorded"])]
        )
        if values.empty:
            return ""
        return "; ".join(values.value_counts().head(4).index.tolist())

    hint = summarize("pair")
    genotype = summarize("genotype")
    return hint, genotype, int(len(frame))


def glob_count(path: Path, pattern: str) -> int:
    return sum(1 for _ in path.glob(pattern)) if path.is_dir() else 0


def stage_state(experiment_id: str) -> tuple[dict[str, int], list[dict], dict]:
    acq = PIPELINE_ROOT / "acquisition" / experiment_id
    obj = PIPELINE_ROOT / "object_extraction" / experiment_id
    qc = PIPELINE_ROOT / "quality_control" / experiment_id
    feat = PIPELINE_ROOT / "feature_extraction" / experiment_id
    ready = PIPELINE_ROOT / "analysis_ready" / experiment_id

    frame_dir = acq / "frame_inventory/per_well"
    frame_files = sorted(frame_dir.glob("*/*_frame_inventory.csv"))
    n_frameshards = len(frame_files)
    n_frame_validated = glob_count(frame_dir, "*/*_frame_inventory.csv.validated")
    target_wells = n_frameshards
    approx_frame_count, frame_basis = approximate_frames(frame_files)

    metadata_files = [
        acq / "ingest_metadata/plate_metadata.csv.validated",
        acq / "ingest_metadata/scope_metadata_mapped.csv.validated",
    ]
    metadata_count = sum(path.is_file() for path in metadata_files)

    per_well_specs = {
        "frame_detections": (
            obj / "frame_detections/per_well",
            "*/*_frame_detections.csv",
        ),
        "frame_masks": (
            obj / "frame_masks/per_well",
            "*/*_frame_masks.csv.validated",
        ),
        "physical_embryo_registry": (
            obj / "physical_embryo_registry/per_well",
            "*/*_physical_embryo_registry.csv.validated",
        ),
        "snip_processing": (
            obj / "snips/per_well",
            "*/*_snip_inventory.csv.validated",
        ),
        "focus_qc": (qc / "focus_qc/per_well", "*/*_focus_qc.csv.validated"),
        "motion_blur_qc": (
            qc / "motion_blur_qc/per_well",
            "*/*_motion_blur_qc.csv.validated",
        ),
        "mask_quality_qc": (
            qc / "mask_quality_qc/per_well",
            "*/*_mask_quality_qc.csv.validated",
        ),
        "surface_area_qc": (
            qc / "surface_area_qc/per_well",
            "*/*_surface_area_qc.csv.validated",
        ),
    }
    merged_specs = {
        "snip_qc": qc / "snip_qc" / f"{experiment_id}_snip_qc.parquet",
        "death_detection": qc
        / "death_detection"
        / f"{experiment_id}_death_event.csv",
        "mask_geometry": feat
        / "mask_geometry"
        / f"{experiment_id}_mask_geometry.csv",
        "curvature_metrics": feat
        / "curvature_metrics"
        / f"{experiment_id}_curvature_metrics.csv",
        "stage_predictions": feat
        / "stage_predictions"
        / f"{experiment_id}_stage_predictions.csv",
        "latent_embeddings": feat
        / "latent_embeddings"
        / f"{experiment_id}_latents.parquet",
        "pose_kinematics": feat
        / "pose_kinematics"
        / f"{experiment_id}_pose_kinematics.csv",
        "fraction_alive": feat
        / "fraction_alive"
        / f"{experiment_id}_fraction_alive.csv",
        "analysis_ready": ready
        / "analysis_ready"
        / f"{experiment_id}_analysis_ready.parquet",
    }

    states: dict[str, int] = {
        "metadata_ingest": int(
            metadata_count == len(metadata_files)
            or (n_frameshards > 0 and n_frame_validated == n_frameshards)
        ),
        "materialization": int(
            n_frameshards > 0 and n_frame_validated == n_frameshards
        ),
    }
    evidence = [
        {
            "experiment_id": experiment_id,
            "stage": "metadata_ingest",
            "stage_label": dict(STAGES)["metadata_ingest"],
            "complete": states["metadata_ingest"],
            "observed_count": metadata_count,
            "expected_count": len(metadata_files),
            "evidence_path": str(acq / "ingest_metadata"),
        },
        {
            "experiment_id": experiment_id,
            "stage": "materialization",
            "stage_label": dict(STAGES)["materialization"],
            "complete": states["materialization"],
            "observed_count": n_frame_validated,
            "expected_count": n_frameshards,
            "evidence_path": str(frame_dir),
        },
    ]
    for stage, (directory, pattern) in per_well_specs.items():
        observed = glob_count(directory, pattern)
        complete = int(target_wells > 0 and observed >= target_wells)
        states[stage] = complete
        evidence.append(
            {
                "experiment_id": experiment_id,
                "stage": stage,
                "stage_label": dict(STAGES)[stage],
                "complete": complete,
                "observed_count": observed,
                "expected_count": target_wells,
                "evidence_path": str(directory),
            }
        )
    for stage, path in merged_specs.items():
        complete = int(path.is_file())
        states[stage] = complete
        evidence.append(
            {
                "experiment_id": experiment_id,
                "stage": stage,
                "stage_label": dict(STAGES)[stage],
                "complete": complete,
                "observed_count": complete,
                "expected_count": 1,
                "evidence_path": str(path),
            }
        )
    extra = {
        "well_count": target_wells or None,
        "validated_well_count": n_frame_validated,
        "approx_frame_count": approx_frame_count,
        "approx_frame_count_basis": frame_basis,
    }
    return states, evidence, extra


def infer_failure_from_error(signature: str) -> Failure | None:
    low = signature.lower()
    if not low:
        return None
    rules = [
        (
            "missing promised flag column",
            _failure(
                "Artifact compatibility",
                "Pre-refactor QC table has a stale schema",
                "focus_qc",
                1,
                "rerun_needed",
                signature,
                "Invalidate and rebuild the affected QC products under the current schema.",
            ),
        ),
        (
            "permissionerror",
            _failure(
                "Filesystem permissions",
                "Work-directory permission error",
                "frame_detections",
                1,
                "operational_cleanup",
                signature,
                "Use a writable work-directory root or repair ownership, then requeue.",
            ),
        ),
        (
            "emptydataerror",
            _failure(
                "Detection/segmentation",
                "Headerless empty intermediate CSV",
                "snip_processing",
                3,
                "segmentation_retry_needed",
                signature,
                "Emit schemaful empty output, quarantine the well, and retry segmentation.",
            ),
        ),
        (
            "death_event: no stage_predictions",
            _failure(
                "Join semantics",
                "Missing death-stage join row",
                "death_detection",
                2,
                "code_change_needed",
                signature,
                "Calculate death stage from well/time metadata rather than embryo/snips.",
            ),
        ),
        (
            "aligned mask contains zero pixels",
            _failure(
                "QC robustness",
                "Zero-pixel aligned mask",
                "motion_blur_qc",
                3,
                "well_quarantine_needed",
                signature,
                "Mark QC not applicable, quarantine the mask, and retry segmentation.",
            ),
        ),
        (
            "stale snakemake lock",
            _failure(
                "Workflow state",
                "Stale Snakemake lock",
                "frame_detections",
                1,
                "operational_cleanup",
                signature,
                "Unlock the inactive work directory and requeue.",
            ),
        ),
    ]
    for token, issue in rules:
        if token in low:
            return issue
    return _failure(
        "Unclassified runtime failure",
        "Requires focused log review",
        "analysis_ready",
        4,
        "diagnosis_needed",
        signature,
        "Inspect the referenced latest stderr log and assign a stable failure category.",
    )


def audit_datasets() -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    universe = parse_manifest()
    logs = collect_logs()
    known = known_failures()
    audit_time = datetime.now().astimezone()
    rows = []
    evidence_rows: list[dict] = []

    for record in universe.to_dict("records"):
        experiment_id = record["experiment_id"]
        scope = record["scope"]
        states, evidence, extra = stage_state(experiment_id)
        evidence_rows.extend({**row, "scope": scope} for row in evidence)
        front_log = latest_log(logs, experiment_id, "front_half")
        back_log = latest_log(logs, experiment_id, "back_half")
        latest = back_log or front_log

        issues = list(known.get(experiment_id, []))
        inferred = infer_failure_from_error(str(latest.get("error_signature", "")))
        if (
            inferred
            and inferred.category not in {issue.category for issue in issues}
            and (
                not issues
                or inferred.category != "Unclassified runtime failure"
            )
        ):
            issues.append(inferred)
        issues.sort(key=lambda issue: (STAGE_ORDER.get(issue.stage, 999), issue.effort_tier))

        analysis_present = states["analysis_ready"]
        latest_back_state = str(back_log.get("run_state", "not_attempted"))
        latest_front_state = str(front_log.get("run_state", "not_attempted"))
        confirmed_clean = int(
            analysis_present == 1 and latest_back_state == "success"
        )

        if latest_back_state == "running_or_incomplete":
            pipeline_status = "current_run_in_progress"
        elif confirmed_clean:
            pipeline_status = "complete_confirmed"
        elif analysis_present:
            pipeline_status = "terminal_product_present_not_latest-clean"
        elif not states["materialization"]:
            pipeline_status = "front_half_incomplete"
        elif latest_back_state in {"failed", "incomplete_unknown"}:
            pipeline_status = "back_half_failed"
        else:
            pipeline_status = "back_half_not_attempted"

        if pipeline_status == "back_half_not_attempted" and not issues:
            issues = [
                _failure(
                    "Not previously attempted",
                    "Front half is complete; back half has not run",
                    "frame_detections",
                    1,
                    "ready_to_run",
                    "",
                    "Submit through the standard back-half array; reuse validated front-half "
                    "artifacts.",
                )
            ]

        primary = issues[0] if issues else None
        blocking_tier = max((issue.effort_tier for issue in issues), default=0)
        completed_stages = sum(states.values())
        first_incomplete = next(
            (name for name, _ in STAGES if not states[name]), ""
        )
        project_hint, genotype_hint, plate_rows = metadata_hints(experiment_id)

        date_match = re.match(r"(\d{8})", experiment_id)
        acquisition_date = (
            datetime.strptime(date_match.group(1), "%Y%m%d").date()
            if date_match
            else pd.NaT
        )
        row = {
            "experiment_id": experiment_id,
            "scope": scope,
            "project_name": project_name(experiment_id),
            "metadata_project_hint": project_hint,
            "genotype_hint": genotype_hint,
            "acquisition_date": acquisition_date,
            "pipeline_status": pipeline_status,
            "confirmed_clean_latest_run": confirmed_clean,
            "terminal_product_present": analysis_present,
            "pipeline_completion_fraction": round(completed_stages / len(STAGES), 3),
            "first_incomplete_stage": first_incomplete,
            "well_count": extra["well_count"] or plate_rows,
            "validated_well_count": extra["validated_well_count"],
            "approx_frame_count": extra["approx_frame_count"],
            "approx_frame_count_basis": extra["approx_frame_count_basis"],
            "failure_category": primary.category if primary else "",
            "failure_subcategory": primary.subcategory if primary else "",
            "all_failure_categories": "; ".join(
                dict.fromkeys(issue.category for issue in issues)
            ),
            "blocking_issue_count": len(issues),
            "failure_stage": primary.stage if primary else "",
            "error_signature": (
                primary.error_signature
                if primary and primary.error_signature
                else str(latest.get("error_signature", ""))
            ),
            "latest_run_kind": latest.get("run_kind", ""),
            "latest_job_id": latest.get("job_id", pd.NA),
            "latest_task_id": latest.get("task_id", pd.NA),
            "latest_run_state": latest.get("run_state", "not_attempted"),
            "latest_exit_status": latest.get("inferred_exit_status", pd.NA),
            "latest_log_time": latest.get("log_mtime", pd.NaT),
            "failure_evidence_path": latest.get("stderr_path", ""),
            "partial_outputs_present": int(
                completed_stages > 0 and analysis_present == 0
            ),
            "recommended_solution": primary.solution if primary else "",
            "remediation_state": primary.remediation_state if primary else "none",
            "recovery_tier": blocking_tier,
            "recovery_tier_label": {
                0: "Complete / none",
                1: "1 - Rerun or operational cleanup",
                2: "2 - Small general code patch",
                3: "3 - Targeted data/model recovery",
                4: "4 - Diagnosis required",
            }[blocking_tier],
            "quick_win": int(0 < blocking_tier <= 2),
            "expected_recoverable": int(bool(issues) and blocking_tier <= 3),
            "retry_priority": (
                "P0"
                if blocking_tier == 1
                else "P1"
                if blocking_tier == 2
                else "P2"
                if blocking_tier == 3
                else "P3"
                if blocking_tier == 4
                else ""
            ),
            "audit_timestamp": audit_time.isoformat(),
        }
        for stage, _ in STAGES:
            row[f"stage_{stage}"] = states[stage]
        rows.append(row)

    audit = pd.DataFrame(rows).sort_values(
        ["scope", "confirmed_clean_latest_run", "recovery_tier", "experiment_id"],
        ascending=[True, True, True, True],
    )
    return audit, pd.DataFrame(evidence_rows), logs


def build_failure_details(audit: pd.DataFrame) -> pd.DataFrame:
    known = known_failures()
    rows = []
    for record in audit.to_dict("records"):
        issues = list(known.get(record["experiment_id"], []))
        inferred = infer_failure_from_error(str(record.get("error_signature", "")))
        if (
            inferred
            and inferred.category not in {issue.category for issue in issues}
            and (
                not issues
                or inferred.category != "Unclassified runtime failure"
            )
        ):
            issues.append(inferred)
        if not issues and record.get("failure_category"):
            issues.append(
                _failure(
                    str(record["failure_category"]),
                    str(record["failure_subcategory"]),
                    str(record["failure_stage"]),
                    int(record["recovery_tier"]),
                    str(record["remediation_state"]),
                    str(record["error_signature"]),
                    str(record["recommended_solution"]),
                )
            )
        unique = {
            (issue.category, issue.subcategory): issue for issue in issues
        }
        for issue in unique.values():
            rows.append(
                {
                    "experiment_id": record["experiment_id"],
                    "scope": record["scope"],
                    "project_name": record["project_name"],
                    "currently_blocking": int(
                        record["confirmed_clean_latest_run"] == 0
                    ),
                    "failure_category": issue.category,
                    "failure_subcategory": issue.subcategory,
                    "failure_stage": issue.stage,
                    "recovery_tier": issue.effort_tier,
                    "remediation_state": issue.remediation_state,
                    "error_signature": issue.error_signature,
                    "recommended_solution": issue.solution,
                    "failure_evidence_path": record["failure_evidence_path"],
                }
            )
    return pd.DataFrame(rows).sort_values(
        ["currently_blocking", "recovery_tier", "failure_category", "experiment_id"],
        ascending=[False, True, True, True],
    )


def build_taxonomy(details: pd.DataFrame) -> pd.DataFrame:
    active = details[details["currently_blocking"] == 1].copy()
    active["recovery_tier_label"] = active["recovery_tier"].map(
        {
            1: "1 - Rerun or operational cleanup",
            2: "2 - Small general code patch",
            3: "3 - Targeted data/model recovery",
            4: "4 - Diagnosis required",
        }
    )
    grouped = (
        active.groupby(
            [
                "failure_category",
                "failure_subcategory",
                "failure_stage",
                "recovery_tier",
                "recovery_tier_label",
            ],
            dropna=False,
        )
        .agg(
            remediation_state=(
                "remediation_state",
                lambda values: "; ".join(sorted(set(values))),
            ),
            recommended_solution=("recommended_solution", "first"),
            dataset_count=("experiment_id", "nunique"),
            keyence_count=("scope", lambda s: int((s == "Keyence").sum())),
            yx1_count=("scope", lambda s: int((s == "YX1").sum())),
            datasets=(
                "experiment_id",
                lambda values: ", ".join(sorted(set(values))),
            ),
        )
        .reset_index()
    )
    return grouped.sort_values(
        ["recovery_tier", "dataset_count", "failure_category"],
        ascending=[True, False, True],
    )


def build_recovery(audit: pd.DataFrame) -> pd.DataFrame:
    unresolved = audit[audit["confirmed_clean_latest_run"] == 0].copy()
    rows = []
    for tier in (1, 2, 3, 4):
        selected = unresolved[unresolved["recovery_tier"] == tier]
        rows.append(
            {
                "recovery_tier": tier,
                "recovery_tier_label": {
                    1: "Rerun / operational cleanup",
                    2: "Small general code patch",
                    3: "Targeted data/model recovery",
                    4: "Diagnosis required",
                }[tier],
                "dataset_count": selected["experiment_id"].nunique(),
                "keyence_count": int((selected["scope"] == "Keyence").sum()),
                "yx1_count": int((selected["scope"] == "YX1").sum()),
                "approx_frames_at_stake": int(
                    selected["approx_frame_count"].fillna(0).sum()
                ),
                "datasets": ", ".join(selected["experiment_id"].sort_values()),
            }
        )
    return pd.DataFrame(rows)


def build_summary(audit: pd.DataFrame) -> pd.DataFrame:
    total = len(audit)
    confirmed = int(audit["confirmed_clean_latest_run"].sum())
    terminal = int(audit["terminal_product_present"].sum())
    target = math.ceil(0.90 * total)
    quick = int(
        audit.loc[
            (audit["confirmed_clean_latest_run"] == 0)
            & (audit["recovery_tier"].isin([1, 2])),
            "experiment_id",
        ].nunique()
    )
    known_recoverable = int(
        audit.loc[
            (audit["confirmed_clean_latest_run"] == 0)
            & (audit["recovery_tier"].isin([1, 2, 3])),
            "experiment_id",
        ].nunique()
    )
    return pd.DataFrame(
        [
            ("Audit timestamp", audit["audit_timestamp"].iloc[0]),
            ("Dataset universe", total),
            ("Keyence datasets", int((audit["scope"] == "Keyence").sum())),
            ("YX1 datasets", int((audit["scope"] == "YX1").sum())),
            ("Confirmed clean latest run", confirmed),
            ("Confirmed clean fraction", confirmed / total),
            ("Terminal product present", terminal),
            ("Terminal product fraction", terminal / total),
            ("90% target dataset count", target),
            ("Additional clean datasets needed", max(0, target - confirmed)),
            ("Tier 1-2 unresolved datasets", quick),
            ("Tier 1-3 known-recoverable datasets", known_recoverable),
            (
                "Potential confirmed after Tier 1-3",
                min(total, confirmed + known_recoverable),
            ),
        ],
        columns=["metric", "value"],
    )


def data_dictionary(audit: pd.DataFrame) -> pd.DataFrame:
    descriptions = {
        "experiment_id": "Canonical pipeline experiment/dataset identifier.",
        "scope": "Microscope source: Keyence or YX1.",
        "project_name": "Coarse science-project grouping inferred from experiment name.",
        "pipeline_status": "Strict status combining terminal artifacts and latest run state.",
        "confirmed_clean_latest_run": "1 only when analysis_ready exists and the latest back-half log has a FINISHED marker.",
        "terminal_product_present": "1 when the analysis-ready parquet exists, regardless of latest rerun state.",
        "pipeline_completion_fraction": "Mean of the 19 binary stage-completion columns.",
        "approx_frame_count": "Approximate canonical frame-inventory rows, not raw TIFF/ND2 planes.",
        "failure_category": "Primary/earliest classified blocker.",
        "all_failure_categories": "All known blockers; categories can overlap for one dataset.",
        "recovery_tier": "Worst known blocker: 1 operational/rerun, 2 small patch, 3 targeted recovery, 4 diagnosis.",
    }
    rows = []
    for column in audit.columns:
        if column.startswith("stage_"):
            stage = column.removeprefix("stage_")
            description = (
                f"Binary completion for {dict(STAGES).get(stage, stage)}. "
                "Per-well stages require all expected shards; merged stages require the terminal artifact."
            )
        else:
            description = descriptions.get(column, "See column name; retained for detailed audit/reference.")
        rows.append(
            {
                "column": column,
                "description": description,
                "dtype": str(audit[column].dtype),
            }
        )
    return pd.DataFrame(rows)


def format_workbook(path: Path) -> None:
    workbook = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    green = PatternFill("solid", fgColor="C6EFCE")
    red = PatternFill("solid", fgColor="FFC7CE")
    amber = PatternFill("solid", fgColor="FFEB9C")

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for index, column in enumerate(sheet.iter_cols(1, sheet.max_column), start=1):
            values = [str(cell.value or "") for cell in list(column)[:200]]
            width = min(60, max(10, max(map(len, values), default=10) + 2))
            sheet.column_dimensions[get_column_letter(index)].width = width

    audit_sheet = workbook["Dataset Audit"]
    headers = {cell.value: cell.column for cell in audit_sheet[1]}
    for name, column in headers.items():
        if str(name).startswith("stage_") or name in {
            "confirmed_clean_latest_run",
            "terminal_product_present",
            "quick_win",
            "expected_recoverable",
        }:
            letter = get_column_letter(column)
            audit_sheet.conditional_formatting.add(
                f"{letter}2:{letter}{audit_sheet.max_row}",
                CellIsRule(operator="equal", formula=["1"], fill=green),
            )
            audit_sheet.conditional_formatting.add(
                f"{letter}2:{letter}{audit_sheet.max_row}",
                CellIsRule(operator="equal", formula=["0"], fill=red),
            )
    if "recovery_tier" in headers:
        letter = get_column_letter(headers["recovery_tier"])
        audit_sheet.conditional_formatting.add(
            f"{letter}2:{letter}{audit_sheet.max_row}",
            CellIsRule(operator="between", formula=["1", "2"], fill=amber),
        )
    workbook.save(path)


def write_workbook(
    audit: pd.DataFrame,
    evidence: pd.DataFrame,
    logs: pd.DataFrame,
    failure_details: pd.DataFrame,
    taxonomy: pd.DataFrame,
    recovery: pd.DataFrame,
    summary: pd.DataFrame,
) -> None:
    with pd.ExcelWriter(WORKBOOK, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="Executive Summary", index=False)
        audit.to_excel(writer, sheet_name="Dataset Audit", index=False)
        taxonomy.to_excel(writer, sheet_name="Failure Taxonomy", index=False)
        failure_details.to_excel(writer, sheet_name="Failure Details", index=False)
        recovery.to_excel(writer, sheet_name="Recovery Priorities", index=False)
        evidence.to_excel(writer, sheet_name="Stage Evidence", index=False)
        logs.to_excel(writer, sheet_name="Run Log Index", index=False)
        data_dictionary(audit).to_excel(writer, sheet_name="Data Dictionary", index=False)
    format_workbook(WORKBOOK)


def notebook_cells(taxonomy: pd.DataFrame) -> list:
    cells = [
        nbf.v4.new_markdown_cell(
            """# Keyence + YX1 pipeline failure dashboard

This dashboard is generated from `pipeline_audit.xlsx`. It separates:

- **Confirmed clean:** terminal product exists and the latest back-half run finished cleanly.
- **Terminal product present:** a parquet exists, but the latest run may have failed or be unconfirmed.
- **Stage completion:** validated per-well shards or the appropriate merged terminal artifact.

Run `refresh_pipeline_audit.py` to rebuild the workbook and regenerate this notebook."""
        ),
        nbf.v4.new_code_cell(
            """from pathlib import Path
import math
import pandas as pd
import numpy as np
import matplotlib.pyplot as plt
from IPython.display import display, Markdown

plt.style.use("seaborn-v0_8-whitegrid")
WORKBOOK = Path("pipeline_audit.xlsx")
audit = pd.read_excel(WORKBOOK, sheet_name="Dataset Audit")
summary = pd.read_excel(WORKBOOK, sheet_name="Executive Summary")
taxonomy = pd.read_excel(WORKBOOK, sheet_name="Failure Taxonomy")
recovery = pd.read_excel(WORKBOOK, sheet_name="Recovery Priorities")
display(summary)"""
        ),
        nbf.v4.new_markdown_cell("## Executive progress"),
        nbf.v4.new_code_cell(
            """scope = audit.groupby("scope").agg(
    datasets=("experiment_id", "size"),
    confirmed_clean=("confirmed_clean_latest_run", "sum"),
    terminal_present=("terminal_product_present", "sum"),
)
scope["confirmed_fraction"] = scope["confirmed_clean"] / scope["datasets"]
scope["terminal_fraction"] = scope["terminal_present"] / scope["datasets"]
display(scope)

ax = scope[["confirmed_clean", "terminal_present", "datasets"]].plot(
    kind="bar", figsize=(9, 5), color=["#2A9D8F", "#E9C46A", "#D9D9D9"]
)
ax.set_ylabel("Datasets")
ax.set_title("Pipeline coverage by microscope scope")
ax.tick_params(axis="x", rotation=0)
plt.tight_layout()
plt.show()"""
        ),
        nbf.v4.new_markdown_cell("## Stage-completion funnel"),
        nbf.v4.new_code_cell(
            """stage_cols = [c for c in audit.columns if c.startswith("stage_")]
stage_labels = [c.removeprefix("stage_").replace("_", " ").title() for c in stage_cols]
stage_by_scope = audit.groupby("scope")[stage_cols].mean().T
stage_by_scope.index = stage_labels

fig, ax = plt.subplots(figsize=(10, 8))
image = ax.imshow(stage_by_scope.T.values, aspect="auto", vmin=0, vmax=1, cmap="RdYlGn")
ax.set_yticks(range(len(stage_by_scope.columns)), stage_by_scope.columns)
ax.set_xticks(range(len(stage_by_scope.index)), stage_by_scope.index, rotation=65, ha="right")
for y in range(stage_by_scope.shape[1]):
    for x in range(stage_by_scope.shape[0]):
        ax.text(x, y, f"{stage_by_scope.iloc[x, y]:.0%}", ha="center", va="center", fontsize=8)
fig.colorbar(image, ax=ax, label="Fraction complete")
ax.set_title("Stage completion by scope")
plt.tight_layout()
plt.show()"""
        ),
        nbf.v4.new_markdown_cell("## Failure modes"),
        nbf.v4.new_code_cell(
            """failure_counts = (
    audit.loc[audit["confirmed_clean_latest_run"] == 0]
    .groupby(["failure_category", "scope"])
    .size()
    .unstack(fill_value=0)
    .sort_values(by=list(audit["scope"].unique()), ascending=False)
)
display(failure_counts)
ax = failure_counts.plot(kind="barh", stacked=True, figsize=(10, max(5, len(failure_counts) * 0.42)))
ax.set_xlabel("Datasets")
ax.set_ylabel("")
ax.set_title("Primary blocker by scope")
plt.tight_layout()
plt.show()"""
        ),
        nbf.v4.new_markdown_cell("## Path to 90%: easiest wins first"),
        nbf.v4.new_code_cell(
            """total = len(audit)
current = int(audit["confirmed_clean_latest_run"].sum())
target = math.ceil(0.90 * total)
tier_counts = (
    audit.loc[audit["confirmed_clean_latest_run"] == 0]
    .groupby("recovery_tier")["experiment_id"].nunique()
    .reindex([1, 2, 3, 4], fill_value=0)
)
cumulative = [current]
for tier in (1, 2, 3, 4):
    cumulative.append(cumulative[-1] + int(tier_counts[tier]))

labels = ["Current", "+ Tier 1", "+ Tier 2", "+ Tier 3", "+ Tier 4"]
fig, ax = plt.subplots(figsize=(9, 5))
ax.plot(labels, cumulative, marker="o", linewidth=3, color="#264653")
ax.axhline(target, color="#E76F51", linestyle="--", label=f"90% target ({target})")
ax.set_ylim(0, total * 1.05)
ax.set_ylabel("Potential confirmed datasets")
ax.set_title("Cumulative recovery opportunity")
ax.legend()
for i, value in enumerate(cumulative):
    ax.text(i, value + 2, str(value), ha="center")
plt.tight_layout()
plt.show()

easy = audit[
    (audit["confirmed_clean_latest_run"] == 0)
    & (audit["recovery_tier"].isin([1, 2]))
][[
    "experiment_id", "scope", "failure_category", "failure_subcategory",
    "recovery_tier_label", "recommended_solution", "approx_frame_count"
]].sort_values(["recovery_tier_label", "scope", "failure_category", "experiment_id"])
display(Markdown(f"**Tier 1–2 opportunity: {len(easy)} datasets. "
                 f"Additional datasets needed for 90%: {max(0, target-current)}.**"))
display(easy)"""
        ),
        nbf.v4.new_markdown_cell("## Detailed failure-mode reference"),
    ]
    for row in taxonomy.to_dict("records"):
        category = str(row["failure_category"])
        subcategory = str(row["failure_subcategory"])
        solution = str(row["recommended_solution"])
        datasets = str(row["datasets"])
        cells.append(
            nbf.v4.new_markdown_cell(
                f"""### {category}: {subcategory}

**Affected:** {int(row['dataset_count'])} datasets  
**Recovery tier:** {int(row['recovery_tier'])} — {row['recovery_tier_label']}  
**Pipeline stage:** `{row['failure_stage']}`  
**Remediation state:** `{row['remediation_state']}`

**Recommended solution:** {solution}

**Datasets:** {datasets}"""
            )
        )
    cells.extend(
        [
            nbf.v4.new_markdown_cell("## Dataset-level reference"),
            nbf.v4.new_code_cell(
                """reference_columns = [
    "experiment_id", "scope", "project_name", "pipeline_status",
    "pipeline_completion_fraction", "first_incomplete_stage",
    "approx_frame_count", "failure_category", "failure_subcategory",
    "latest_job_id", "latest_task_id", "latest_run_state",
    "recovery_tier_label", "recommended_solution", "failure_evidence_path",
]
display(audit.sort_values(
    ["scope", "confirmed_clean_latest_run", "recovery_tier", "experiment_id"]
)[reference_columns])"""
            ),
        ]
    )
    return cells


def write_notebook(taxonomy: pd.DataFrame) -> None:
    notebook = nbf.v4.new_notebook()
    notebook["metadata"]["kernelspec"] = {
        "display_name": "Python 3",
        "language": "python",
        "name": "python3",
    }
    notebook["metadata"]["language_info"] = {"name": "python", "version": "3"}
    notebook["cells"] = notebook_cells(taxonomy)
    nbf.write(notebook, NOTEBOOK)


def execute_notebook() -> None:
    subprocess.run(
        [
            sys.executable,
            "-m",
            "nbconvert",
            "--to",
            "notebook",
            "--execute",
            "--inplace",
            "--ExecutePreprocessor.timeout=300",
            NOTEBOOK.name,
        ],
        cwd=AUDIT_DIR,
        check=True,
    )


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--no-execute-notebook",
        action="store_true",
        help="Write the workbook and notebook source without executing notebook cells.",
    )
    args = parser.parse_args()

    audit, evidence, logs = audit_datasets()
    failure_details = build_failure_details(audit)
    taxonomy = build_taxonomy(failure_details)
    recovery = build_recovery(audit)
    summary = build_summary(audit)
    write_workbook(
        audit, evidence, logs, failure_details, taxonomy, recovery, summary
    )
    write_notebook(taxonomy)
    if not args.no_execute_notebook:
        execute_notebook()

    print(f"Wrote {WORKBOOK}")
    print(f"Wrote {NOTEBOOK}")
    print(
        f"Datasets={len(audit)} confirmed_clean={int(audit['confirmed_clean_latest_run'].sum())} "
        f"terminal_present={int(audit['terminal_product_present'].sum())}"
    )


if __name__ == "__main__":
    main()
